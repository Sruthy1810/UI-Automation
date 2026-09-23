from openpyxl import load_workbook
from openpyxl.styles import PatternFill


def color_failed_rows(file_path):

    workbook = load_workbook(file_path)

    sheet = workbook.active

    red_fill = PatternFill(
            fill_type="solid",
            fgColor="FF0000"
        )
    
    no_fill = PatternFill(fill_type=None)

    headers = {}

    for cell in sheet[1]:
        headers[cell.value] = cell.column

    status_column = headers.get("Status")

    if not status_column:
        print("Status column not found")
        workbook.close()
        return

    

    # -----------------------------------------
    # Check every employee row
    # -----------------------------------------

    for row in range(2, sheet.max_row + 1):

        status = sheet.cell(
            row=row,
            column=status_column
        ).value

        # -------------------------------------
        # REMOVE OLD COLOR FIRST
        # -------------------------------------

        for column in range(1, sheet.max_column + 1):

            sheet.cell(
                row=row,
                column=column
            ).fill = no_fill

        # -------------------------------------
        # APPLY RED ONLY FOR CURRENT FAILURE
        # -------------------------------------

        if status and str(status).strip().lower() == "failed":

            for column in range(1, sheet.max_column + 1):

                sheet.cell(
                    row=row,
                    column=column
                ).fill = red_fill

    # -----------------------------------------
    # Add filter
    # -----------------------------------------

    sheet.auto_filter.ref = sheet.dimensions

    # -----------------------------------------
    # Save Excel
    # -----------------------------------------

    workbook.save(file_path)

    workbook.close()

    print("Excel formatting updated successfully")