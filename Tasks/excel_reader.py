from openpyxl import load_workbook


def read_employee_data(file_path):

    workbook = load_workbook(file_path)

    sheet = workbook.active

    employees = []

    headers = [
        cell.value
        for cell in sheet[1]
    ]


    for row_number, row in enumerate(
        sheet.iter_rows(min_row=2, values_only=True),
        start=2
    ):

        employee = dict(
            zip(headers, row)
        )

        employee["_row_number"] = row_number
        employees.append(employee)

    return employees


def update_employee_status(file_path, row_number, status, remarks):

    workbook = load_workbook(file_path)
    sheet = workbook.active

    # Find Status and Remarks columns
    headers = {}

    for col in range(1, sheet.max_column + 1):
        header = sheet.cell(row=1, column=col).value
        if header:
            headers[str(header).strip()] = col

    # Update Status
    status_col = headers["Status"]
    remarks_col = headers["Remarks"]

    # Clear old values first
    sheet.cell(row=row_number, column=status_col).value = None
    sheet.cell(row=row_number, column=remarks_col).value = None

    # Add new values
    sheet.cell(row=row_number, column=status_col).value = status
    sheet.cell(row=row_number, column=remarks_col).value = remarks

    workbook.save(file_path)