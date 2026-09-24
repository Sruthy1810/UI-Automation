from openpyxl import load_workbook


def read_employee_data(file_path):

    workbook = load_workbook(
        file_path
    )

    sheet = workbook.active

    headers = [
        cell.value
        for cell in sheet[1]
    ]

    employees = []

    for row_number, row in enumerate(
        sheet.iter_rows(
            min_row=2,
            values_only=True
        ),
        start=2
    ):

        if not any(row):
            continue

        employee = dict(
            zip(headers, row)
        )

        employee["_row_number"] = row_number

        employees.append(employee)

    workbook.close()

    return employees

def reset_employee_status(file_path):

    workbook = load_workbook(file_path)
    sheet = workbook.active

    headers = {cell.value: cell.column
        for cell in sheet[1]
    }

    status_column = headers.get("Status")
    remarks_column = headers.get("Remarks")

    for row in range(2, sheet.max_row + 1):

        if status_column:
            sheet.cell(
                row=row,
                column=status_column
            ).value = ""

        if remarks_column:
            sheet.cell(
                row=row,
                column=remarks_column
            ).value = ""

    workbook.save(file_path)
    workbook.close()

    print("Previous Status and Remarks cleared")

# ---------------------------------------------------------
# Update Employee Status
# ---------------------------------------------------------


def update_employee_status(
    file_path,
    row_number,
    status,
    remarks
):

    workbook = load_workbook(file_path)

    sheet = workbook.active

    headers = {
        cell.value: cell.column
        for cell in sheet[1]
    }

    # -----------------------------------------
    # Status column
    # -----------------------------------------

    status_column = headers.get("Status")

    if status_column is None:
            status_column = sheet.max_column + 1
            sheet.cell(
                row=1,
                column=status_column
            ).value = "Status"

    # -----------------------------------------
    # Remarks column
    # -----------------------------------------


    remarks_column = headers.get("Remarks")

    if remarks_column is None:
        remarks_column = sheet.max_column + 1
        sheet.cell(
            row=1,
            column=remarks_column
        ).value = "Remarks"


    # -----------------------------------------
    # Update values
    # -----------------------------------------   

    sheet.cell(
        row=row_number,
        column=status_column
    ).value = status

    sheet.cell(
        row=row_number,
        column=remarks_column
    ).value = remarks

    workbook.save(file_path)

    workbook.close()