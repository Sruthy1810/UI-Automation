
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from config.config_reader import get_pdf_files, get_output_path
from modules.pdf_reader import read_scanned_pdf
from modules.ocr_extractor import extract_text_from_pages
from modules.field_extractor import extract_fields


def format_sheet(ws):
    # Format header
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="4472C4"
        )

    # Adjust column widths
    for column in ws.columns:
        letter = column[0].column_letter
        max_length = max(
            len(str(cell.value or ""))
            for cell in column
        )
        ws.column_dimensions[letter].width = min(max_length + 3, 35)

    # Wrap cell content
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main():
    pdf_files = get_pdf_files()

    if not pdf_files:
        print("No PDF files found in the Input folder.")
        return

    wb = Workbook()
    # Remove the default sheet
    default_sheet = wb.active
    wb.remove(default_sheet)

    headers = [
        "Page Number",
        "Title",
        "Account Number",
        "Principal",
        "Review Required"
    ]

    for pdf_path in pdf_files:
        print(f"Processing: {pdf_path.name}")

        # Excel sheet names cannot exceed 31 characters
        sheet_name = pdf_path.stem[:31]

        # Avoid duplicate sheet names
        base_name = sheet_name
        counter = 1
        while sheet_name in wb.sheetnames:
            suffix = f"_{counter}"
            sheet_name = base_name[:31 - len(suffix)] + suffix
            counter += 1

        ws = wb.create_sheet(title=sheet_name)
        ws.append(headers)

        try:
            pages = read_scanned_pdf(str(pdf_path))

            for page_number, page in enumerate(pages, start=1):
                text = extract_text_from_pages([page])
                records = extract_fields(text)

                for record in records:
                    ws.append([
                        page_number,
                        record.get("Title") or "Not found",
                        record.get("Account Number") or "Not found",
                        record.get("Principal") or "Not found",
                        record.get("Review Required", "")
                    ])

        except Exception as error:
            ws.append([
                "",
                "PDF processing failed",
                "",
                "",
                str(error)
            ])
            print(f"Failed: {pdf_path.name}: {error}")

        format_sheet(ws)

    output_path = get_output_path()
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb.save(output_path)

    print(f"Processed {len(pdf_files)} PDF files.")
    print(f"Excel saved: {output_path}")


if __name__ == "__main__":
    main()